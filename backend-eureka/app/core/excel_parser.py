"""
Parser de Excel para importación de plantillas de inspección.

Mejoras v2:
  - Detecta subencabezados (celdas fusionadas / doble fila de encabezados)
  - Maneja "Fecha de recarga → Actual / Final" como dos campos separados
  - Mejor detección de FORMULARIO vs MATRIZ vs FORMULARIO_MATRIZ
"""

import re
from typing import Any, Dict, List, Optional

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_KW_DATE    = {"fecha", "date", "vencimiento", "vigencia", "plazo", "caducidad", "recarga"}
_KW_CHECK_SN= {"cumple", "sí", "si", "no", "aplica", "n/a", "na", "conforme",
               "ok", "check", "cumplimiento"}
_KW_CHECK_BM= {"bueno", "malo", "regular", "estado", "condición", "condicion"}
_KW_OBS     = {"observaci", "nota", "comentario", "descripci", "detalle", "remark"}
_KW_FOTO    = {"foto", "imagen", "photo", "evidencia", "image"}
_KW_NUM     = {"número", "numero", "cantidad", "capacidad", "peso", "kg", "litro",
               "presión", "presion", "temperatura", "psi", "bar", "code", "código"}
_KW_GENERAL = {"empresa", "organización", "organizacion", "area", "área",
               "responsable", "inspector", "fecha", "turno", "período", "periodo",
               "sede", "sucursal", "departamento", "cargo", "jefe"}
_KW_ANSWER  = {"sí", "si", "no", "n/a", "na", "x", "✓", "✗", "ok",
               "bueno", "malo", "conforme", "no conforme"}


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()

def _slugify(text: str) -> str:
    s = _normalize(text)
    s = re.sub(r"[^a-z0-9\s]", "", s)
    s = re.sub(r"\s+", "_", s).strip("_")
    return s[:50] or "campo"

def _cell_value(cell) -> str:
    if cell is None or cell.value is None:
        return ""
    return str(cell.value).strip()

def _is_empty_row(row) -> bool:
    return all(_cell_value(c) == "" for c in row)

def _detect_field_type(header: str, sample_values: List[str]) -> str:
    h = _normalize(header)
    if any(kw in h for kw in _KW_DATE):     return "fecha"
    if any(kw in h for kw in _KW_OBS):      return "observacion"
    if any(kw in h for kw in _KW_FOTO):     return "foto"
    if any(kw in h for kw in _KW_NUM):      return "numero"
    sv = {_normalize(v) for v in sample_values if v}
    if sv and sv <= _KW_ANSWER:
        if {"bueno", "malo", "regular"} & sv: return "check_bm"
        if {"n/a", "na"} & sv:                return "check_sna"
        return "check_sn"
    if any(kw in h for kw in _KW_CHECK_SN): return "check_sna"
    if any(kw in h for kw in _KW_CHECK_BM): return "check_bm"
    return "texto"


def _get_merged_ranges(ws):
    """Retorna un dict {(row,col): texto_del_encabezado_fusionado}."""
    merged = {}
    for mr in ws.merged_cells.ranges:
        cell = ws.cell(mr.min_row, mr.min_col)
        val  = _cell_value(cell)
        if val:
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    merged[(r, c)] = val
    return merged


def _detect_header_rows(ws, all_rows, merged):
    """
    Detecta si hay 1 o 2 filas de encabezado.
    Retorna (header_row_idx, subheader_row_idx_or_None, pre_header_rows).

    Caso 1 fila:  | Código | Ubicación | Fecha | Estado |
    Caso 2 filas: | Fecha de recarga       | Estado |
                  | Actual  | Final       |        |
    """
    header_row_idx    = None
    subheader_row_idx = None
    pre_header_rows   = []

    for i, row in enumerate(all_rows):
        if _is_empty_row(row):
            continue
        non_empty = [_cell_value(c) for c in row if _cell_value(c)]
        if len(non_empty) >= 3:
            header_row_idx = i
            # Verificar si la siguiente fila no vacía también es un encabezado
            for j in range(i + 1, min(i + 3, len(all_rows))):
                next_row = all_rows[j]
                if _is_empty_row(next_row):
                    continue
                next_vals = [_cell_value(c) for c in next_row if _cell_value(c)]
                # Si la siguiente fila tiene valores cortos y no parecen datos → subencabezado
                if len(next_vals) >= 2:
                    avg_len = sum(len(v) for v in next_vals) / len(next_vals)
                    looks_like_data = any(
                        v.replace(".", "").replace(",", "").isdigit() or
                        len(v) > 30
                        for v in next_vals
                    )
                    if avg_len < 20 and not looks_like_data:
                        subheader_row_idx = j
                break
            break
        else:
            pre_header_rows.append((i, row))

    return header_row_idx, subheader_row_idx, pre_header_rows


def _build_headers(ws, all_rows, header_row_idx, subheader_row_idx, merged):
    """
    Construye la lista de encabezados finales manejando celdas fusionadas
    y doble fila de encabezados.

    Caso simple:   ["Código", "Ubicación", "Fecha vencimiento", "Estado"]
    Caso doble:    ["Código", "Ubicación", "Fecha recarga - Actual",
                   "Fecha recarga - Final", "Estado"]
    """
    if header_row_idx is None:
        return [], {}

    header_row = all_rows[header_row_idx]
    h_row_num  = header_row_idx + 1  # 1-indexed para openpyxl

    if subheader_row_idx is None:
        # Una sola fila de encabezados
        headers = []
        col_map = {}  # col_idx → header_name
        for ci, cell in enumerate(header_row):
            val = merged.get((h_row_num, ci + 1)) or _cell_value(cell)
            if val:
                headers.append(val)
                col_map[ci] = val
        return headers, col_map

    # Dos filas de encabezados
    sub_row     = all_rows[subheader_row_idx]
    sub_row_num = subheader_row_idx + 1
    headers     = []
    col_map     = {}

    # Determinar qué columnas tienen subencabezado
    parent_by_col = {}  # ci → texto del padre (fila 1)
    for ci, cell in enumerate(header_row):
        val = merged.get((h_row_num, ci + 1)) or _cell_value(cell)
        if val:
            parent_by_col[ci] = val

    for ci, sub_cell in enumerate(sub_row):
        sub_val = _cell_value(sub_cell)
        parent  = parent_by_col.get(ci, "")

        if sub_val and parent:
            # Combinación: "Fecha recarga - Actual"
            full = f"{parent} - {sub_val}"
        elif sub_val and not parent:
            full = sub_val
        elif parent and not sub_val:
            # El padre ocupa esta columna sin subencabezado
            full = parent
        else:
            continue

        headers.append(full)
        col_map[ci] = full

    return headers, col_map


def parse_excel(file_bytes: bytes) -> Dict[str, Any]:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl no está instalado")

    import io
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=False, data_only=True)
    ws = wb.active

    merged    = _get_merged_ranges(ws)
    all_rows  = list(ws.iter_rows())
    warnings  = []

    if not all_rows:
        return _empty_result("El archivo está vacío")

    header_row_idx, subheader_row_idx, pre_header_rows = _detect_header_rows(
        ws, all_rows, merged
    )

    if header_row_idx is None:
        return _empty_result("No se detectó estructura tabular")

    raw_headers, col_map = _build_headers(
        ws, all_rows, header_row_idx, subheader_row_idx, merged
    )

    if not raw_headers:
        return _empty_result("No se detectaron encabezados")

    if subheader_row_idx is not None:
        warnings.append(
            f"Se detectaron encabezados dobles (fila {header_row_idx+1} y {subheader_row_idx+1}). "
            "Los subencabezados se combinaron con sus padres."
        )

    # ── Leer filas de datos ───────────────────────────────────────────────────
    data_start = (subheader_row_idx or header_row_idx) + 1
    data_rows  = []
    for row in all_rows[data_start:data_start + 50]:
        if _is_empty_row(row):
            continue
        vals = [_cell_value(c) for c in row]
        # Alinear usando col_map
        aligned = []
        for ci in sorted(col_map.keys()):
            aligned.append(vals[ci] if ci < len(vals) else "")
        if any(aligned):
            data_rows.append(aligned)

    # ── Muestras por columna ──────────────────────────────────────────────────
    sample_by_col = {h: [] for h in raw_headers}
    for row_vals in data_rows[:10]:
        for h, v in zip(raw_headers, row_vals):
            if v:
                sample_by_col[h].append(v)

    # ── Construir matrix_fields ───────────────────────────────────────────────
    matrix_fields = []
    seen_keys     = set()
    for order, header in enumerate(raw_headers):
        if not header:
            continue
        key = _slugify(header)
        orig = key; c = 1
        while key in seen_keys:
            key = f"{orig}_{c}"; c += 1
        seen_keys.add(key)

        ft = _detect_field_type(header, sample_by_col.get(header, []))
        matrix_fields.append({
            "name":        header,
            "field_key":   key,
            "field_type":  ft,
            "options":     None,
            "is_required": False,
            "order":       order,
            "group_name":  None,
            "scope":       "matriz",
        })

    # ── Detectar campos generales desde filas pre-tabla ───────────────────────
    general_fields = []
    for _, row in pre_header_rows:
        cells     = [_cell_value(c) for c in row]
        non_empty = [c for c in cells if c]
        if len(non_empty) == 2:
            label = non_empty[0].rstrip(":").strip()
            if any(kw in _normalize(label) for kw in _KW_GENERAL) or len(label) < 30:
                ft = _detect_field_type(label, [])
                general_fields.append({
                    "name":        label,
                    "field_key":   _slugify(label),
                    "field_type":  ft,
                    "options":     None,
                    "is_required": False,
                    "order":       len(general_fields),
                    "group_name":  None,
                    "scope":       "general",
                })

    # ── Detectar si es checklist ──────────────────────────────────────────────
    is_checklist = False
    if len(raw_headers) <= 4:
        h_norm = [_normalize(h) for h in raw_headers]
        has_item   = any(any(kw in h for kw in
            {"aspecto", "ítem", "item", "descripci", "verificar", "criterio", "elemento"})
            for h in h_norm)
        has_answer = any(any(kw in h for kw in _KW_CHECK_SN | _KW_CHECK_BM)
            for h in h_norm)
        if has_item and has_answer and len(data_rows) > 2:
            is_checklist = True

    if is_checklist:
        item_col_idx = 0
        for ci, h in enumerate(raw_headers):
            if any(kw in _normalize(h) for kw in
                   {"aspecto", "ítem", "item", "descripci", "verificar", "criterio"}):
                item_col_idx = ci
                break

        other_headers = [h for i, h in enumerate(raw_headers) if i != item_col_idx]
        ft = "check_sna" if any(
            any(kw in _normalize(h) for kw in {"n/a", "na", "aplica"})
            for h in other_headers
        ) else "check_sn"

        checklist_fields = []
        for idx, row_vals in enumerate(data_rows):
            item_name = row_vals[item_col_idx] if row_vals else f"Ítem {idx+1}"
            if not item_name.strip():
                continue
            checklist_fields.append({
                "name":        item_name.strip().rstrip("?"),
                "field_key":   _slugify(item_name),
                "field_type":  ft,
                "options":     None,
                "is_required": False,
                "order":       idx,
                "group_name":  None,
                "scope":       "general",
            })

        if any(any(kw in _normalize(h) for kw in _KW_OBS) for h in raw_headers):
            checklist_fields.append({
                "name": "Observaciones", "field_key": "observaciones",
                "field_type": "observacion", "options": None,
                "is_required": False, "order": len(checklist_fields),
                "group_name": None, "scope": "general",
            })

        general_fields = checklist_fields
        matrix_fields  = []

    # ── Estructura y confianza ────────────────────────────────────────────────
    if is_checklist:
        structure  = "formulario"
        confidence = "alta" if len(general_fields) >= 3 else "media"
    elif general_fields and matrix_fields:
        structure  = "formulario_matriz"
        confidence = "alta"
    elif matrix_fields:
        structure  = "matriz"
        confidence = "alta" if len(data_rows) > 1 else "media"
    else:
        structure  = "formulario"
        confidence = "baja"
        warnings.append("No se pudo determinar la estructura con certeza")

    # ── Preview ───────────────────────────────────────────────────────────────
    preview_rows = [dict(zip(raw_headers, r)) for r in data_rows[:5]]

    wb.close()

    return {
        "detected_structure": structure,
        "confidence":         confidence,
        "sheet_name":         ws.title,
        "general_fields":     general_fields,
        "matrix_fields":      matrix_fields,
        "preview_rows":       preview_rows,
        "warnings":           warnings,
        "raw_headers":        raw_headers,
        "total_data_rows":    len(data_rows),
    }


def _empty_result(warning: str) -> Dict[str, Any]:
    return {
        "detected_structure": "formulario", "confidence": "baja",
        "sheet_name": "", "general_fields": [], "matrix_fields": [],
        "preview_rows": [], "warnings": [warning],
        "raw_headers": [], "total_data_rows": 0,
    }
